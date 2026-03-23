#!/usr/bin/env python3
"""
Generador Maestro de Reportes por Ruta — v2.0
Genera un PDF por ruta con datos del día anterior.
Entradas: carpeta entradas/ (xlsx + pdf)
Salidas:  carpeta salidas/  (PDFs con fecha)
"""

import os
import re
import sys
import unicodedata
import pandas as pd
from datetime import datetime, timedelta
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.platypus import Table, TableStyle
import PyPDF2
import openpyxl

# ============================================================
# CONFIGURACIÓN
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CARPETA_ENTRADAS = os.path.join(BASE_DIR, "entradas")
CARPETA_SALIDAS  = os.path.join(BASE_DIR, "salidas")

MIS_RUTAS = [
    "01.01. JUAN SP",
    "01.02. FRANCISCO MORATO",
    "01.03. DIADEMA",
    "01.04. RIBEIRAO PRETO",
    "01.05. PINDA",
    "01.06. RIO",
    "01.07. OSTRAS",
    "01.08. CAMPINAS",
    "02.05. SANTOS",
    "02.06. GUARULHOS",
    "02.07. CABO FRIO CRISTIAN",
    "02.08. SÃO PEDRO",
    "02.09. CANTI",
]

# Colores del diseño
COLOR_HEADER     = colors.HexColor('#1B3A5C')
COLOR_SUBHEADER  = colors.HexColor('#2C5F8A')
COLOR_ACCENT     = colors.HexColor('#E8F0FE')
COLOR_FILA_ALT   = colors.HexColor('#F5F8FC')
COLOR_TEXTO      = colors.HexColor('#1A1A1A')
COLOR_GRIS       = colors.HexColor('#666666')
COLOR_ROJO       = colors.HexColor('#C0392B')
COLOR_VERDE      = colors.HexColor('#27AE60')

# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def _norm(txt: str) -> str:
    """Normaliza texto removiendo acentos y caracteres especiales."""
    if txt is None:
        return ""
    s = str(txt)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def _find_col(df, candidates):
    """Busca columna que coincida con alguno de los candidatos."""
    if df is None or df.empty:
        return None
    norm_map = {_norm(c): c for c in df.columns}
    for cand in candidates:
        n = _norm(cand)
        if n in norm_map:
            return norm_map[n]
    for key, original in norm_map.items():
        for cand in candidates:
            if _norm(cand) in key:
                return original
    return None


def _fmt_dinero(valor) -> str:
    """Formatea un valor numérico como dinero."""
    try:
        v = float(valor)
        return f"${v:,.2f}"
    except (ValueError, TypeError):
        return "$0.00"


def encontrar_archivos(patron_inicio, extension=".xlsx"):
    """Encuentra TODOS los archivos que coincidan con el patrón."""
    if not os.path.isdir(CARPETA_ENTRADAS):
        return []
    return [
        os.path.join(CARPETA_ENTRADAS, f)
        for f in sorted(os.listdir(CARPETA_ENTRADAS))
        if patron_inicio in f and f.endswith(extension)
    ]


# ============================================================
# FUNCIONES DE LECTURA DE DATOS
# ============================================================

def leer_cuadre_por_rutas(ruta_archivo):
    """Lee el archivo de Cuadre con formato especial (Ruta: ...)."""
    if not ruta_archivo:
        return {}
    try:
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
        sheet = wb.active
        datos = {}

        for row_idx in range(1, sheet.max_row + 1):
            cell_val = sheet.cell(row=row_idx, column=1).value
            if cell_val and isinstance(cell_val, str) and cell_val.startswith("Ruta:"):
                nombre = cell_val.replace("Ruta:", "").strip()
                headers = [c.value for c in sheet[row_idx + 1] if c.value]
                data_row = sheet[row_idx + 2]

                ruta_data = {}
                for i, h in enumerate(headers):
                    val = data_row[i].value
                    if isinstance(val, str) and '$' in val:
                        limpio = val.replace('$', '').replace('.', '').replace(',', '.').strip()
                        ruta_data[h] = pd.to_numeric(limpio, errors='coerce')
                    else:
                        ruta_data[h] = val
                datos[nombre] = ruta_data
        return datos
    except Exception as e:
        print(f"  ⚠ Error leyendo Cuadre: {e}")
        return {}


def leer_excel_tabla(ruta_archivo):
    """Lee archivo Excel con formato de tabla (header en fila 9)."""
    if not ruta_archivo:
        return None
    try:
        df = pd.read_excel(ruta_archivo, engine="openpyxl", header=8)
        # Convertir columnas que parecen numéricas
        for col in df.columns:
            if df[col].dtype == "object":
                muestra = df[col].dropna().astype(str).str.strip()
                if muestra.empty:
                    continue
                if muestra.str.match(r'^[\s\$R]*-?[\d\.,]+[\s]*$').any():
                    limpia = (df[col].astype(str)
                              .str.replace(r'R\$\s*', '', regex=True)
                              .str.replace('$', '', regex=False)
                              .str.replace(' ', '', regex=False)
                              .str.replace('.', '', regex=False)
                              .str.replace(',', '.', regex=False)
                              .str.strip())
                    df[col] = pd.to_numeric(limpia, errors='coerce').fillna(0)
        return df
    except Exception as e:
        print(f"  ⚠ Error: {e}")
        return None


def extraer_clientes_pdf(ruta_pdf):
    """Extrae clientes sin pago del PDF."""
    if not ruta_pdf:
        return {}
    try:
        with open(ruta_pdf, "rb") as f:
            pdf = PyPDF2.PdfReader(f)
            texto = "\n".join(page.extract_text() or "" for page in pdf.pages)

        clientes_por_ruta = {}
        ruta_actual = None

        for linea in texto.split("\n"):
            linea = linea.strip()
            if linea.startswith("RUTA:"):
                ruta_actual = linea.replace("RUTA:", "").strip()
                clientes_por_ruta.setdefault(ruta_actual, [])
                continue

            if ruta_actual:
                if any(x in linea.upper() for x in [
                    "COD. CRÉDITO", "CLIENTE", "DÍAS DE ATRASO",
                    "EXCEPCIÓN", "ATRASO", "PENALIZACIÓN", "TOTAL"
                ]):
                    continue

                patron = re.match(
                    r'^\d{1,2}\s+\d{5,6}\s+(.+?)\s+\d+\s+\d+\s+\d+\s+\d+\s*$',
                    linea
                )
                if patron:
                    partes = linea.split()
                    if len(partes) >= 6:
                        cod = partes[1]
                        nombre = ' '.join(partes[2:-4])
                        dias_atraso = int(partes[-4])
                        clientes_por_ruta[ruta_actual].append({
                            'cod': cod,
                            'nombre': nombre,
                            'dias_atraso': dias_atraso
                        })

        return clientes_por_ruta
    except Exception as e:
        print(f"  ⚠ Error PDF: {e}")
        return {}


def filtrar_por_ruta(df, ruta_val):
    """Filtra DataFrame por valor de ruta."""
    if df is None or df.empty:
        return pd.DataFrame()
    col_ruta = _find_col(df, ["RUTA"])
    if not col_ruta:
        return pd.DataFrame()
    mask = df[col_ruta].astype(str).str.strip() == str(ruta_val).strip()
    return df[mask].copy()


# ============================================================
# GENERADOR DE PDF MEJORADO
# ============================================================

class ReportePDF:
    """Genera un PDF profesional con paginación automática."""

    MARGEN = 40
    FOOTER_H = 30

    def __init__(self, path, nombre_ruta, fecha, paginas_totales_estimadas=1):
        self.path = path
        self.nombre_ruta = nombre_ruta
        self.fecha = fecha
        self.ancho, self.alto = letter
        self.c = canvas.Canvas(path, pagesize=letter)
        self.pagina = 1
        self.y = self.alto - 80
        self._dibujar_header()

    @property
    def _espacio_disponible(self):
        return self.y - self.MARGEN - self.FOOTER_H

    def _nueva_pagina(self):
        """Crea una nueva página manteniendo header/footer."""
        self._dibujar_footer()
        self.c.showPage()
        self.pagina += 1
        self.y = self.alto - 80
        self._dibujar_header()

    def _verificar_espacio(self, necesario=60):
        """Si no hay espacio suficiente, salta a nueva página."""
        if self._espacio_disponible < necesario:
            self._nueva_pagina()

    def _dibujar_header(self):
        """Dibuja el encabezado del reporte."""
        c = self.c
        # Barra superior
        c.setFillColor(COLOR_HEADER)
        c.rect(0, self.alto - 65, self.ancho, 65, fill=True, stroke=False)
        # Línea de acento
        c.setFillColor(COLOR_SUBHEADER)
        c.rect(0, self.alto - 70, self.ancho, 5, fill=True, stroke=False)

        # Texto del header
        c.setFillColor(colors.white)
        c.setFont('Helvetica-Bold', 16)
        c.drawString(self.MARGEN, self.alto - 30, f"▸ {self.nombre_ruta}")
        c.setFont('Helvetica', 10)
        c.drawString(self.MARGEN, self.alto - 50, "INFORME DIARIO POR RUTA")

        c.setFont('Helvetica-Bold', 11)
        c.drawRightString(self.ancho - self.MARGEN, self.alto - 30, self.fecha)
        c.setFont('Helvetica', 8)
        c.drawRightString(self.ancho - self.MARGEN, self.alto - 50, "Generado automáticamente")

        self.y = self.alto - 95

    def _dibujar_footer(self):
        """Dibuja el pie de página."""
        c = self.c
        y_footer = 18
        c.setStrokeColor(COLOR_GRIS)
        c.setLineWidth(0.5)
        c.line(self.MARGEN, y_footer + 10, self.ancho - self.MARGEN, y_footer + 10)
        c.setFillColor(COLOR_GRIS)
        c.setFont('Helvetica', 7)
        c.drawString(self.MARGEN, y_footer, f"Reporte generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        c.drawRightString(self.ancho - self.MARGEN, y_footer, f"Página {self.pagina}")

    def seccion(self, titulo):
        """Dibuja un título de sección con línea."""
        self._verificar_espacio(50)
        c = self.c
        # Rectángulo de fondo para el título
        c.setFillColor(COLOR_ACCENT)
        c.rect(self.MARGEN, self.y - 5, self.ancho - 2 * self.MARGEN, 20,
               fill=True, stroke=False)
        c.setFillColor(COLOR_HEADER)
        c.setFont('Helvetica-Bold', 11)
        c.drawString(self.MARGEN + 8, self.y + 1, titulo)

        # Línea inferior
        c.setStrokeColor(COLOR_SUBHEADER)
        c.setLineWidth(1.5)
        c.line(self.MARGEN, self.y - 5, self.ancho - self.MARGEN, self.y - 5)
        self.y -= 25

    def cuadre(self, datos):
        """Dibuja la sección de Cuadre con 3 columnas."""
        self.seccion("CUADRE DE RUTA")
        c = self.c

        abonos = datos.get('ABONOS', 0) or 0
        desembolsos = datos.get('DESEMBOLSOS', 0) or 0
        total_caja = datos.get('TOTAL CAJA', 0) or 0

        col_w = (self.ancho - 2 * self.MARGEN) / 3
        items = [
            ("ABONOS", abonos, COLOR_VERDE),
            ("DESEMBOLSOS", desembolsos, COLOR_ROJO),
            ("TOTAL CAJA", total_caja, COLOR_HEADER),
        ]

        for i, (label, valor, color) in enumerate(items):
            x_centro = self.MARGEN + col_w * i + col_w / 2

            # Label
            c.setFillColor(COLOR_GRIS)
            c.setFont('Helvetica-Bold', 8)
            c.drawCentredString(x_centro, self.y, label)

            # Valor
            c.setFillColor(color)
            c.setFont('Helvetica-Bold', 13)
            c.drawCentredString(x_centro, self.y - 18, _fmt_dinero(valor))

        self.y -= 45

    def tabla(self, titulo, df, cols_config):
        """
        Dibuja una tabla con paginación automática.
        cols_config: lista de tuplas (nombre_busqueda_candidates, label, ancho_pct, alineacion)
        """
        self.seccion(titulo)

        if df is None or df.empty:
            self.c.setFillColor(COLOR_GRIS)
            self.c.setFont('Helvetica-Oblique', 9)
            self.c.drawString(self.MARGEN + 10, self.y, f"Sin datos de {titulo.lower()}")
            self.y -= 20
            return

        # Buscar columnas
        cols_encontradas = []
        labels = []
        anchos = []
        aligns = []
        for candidates, label, pct, align in cols_config:
            col = _find_col(df, candidates)
            if col:
                cols_encontradas.append(col)
                labels.append(label)
                anchos.append((self.ancho - 2 * self.MARGEN) * pct)
                aligns.append(align)

        if not cols_encontradas:
            self.c.setFillColor(COLOR_GRIS)
            self.c.setFont('Helvetica-Oblique', 9)
            self.c.drawString(self.MARGEN + 10, self.y, f"Columnas no encontradas para {titulo}")
            self.y -= 20
            return

        tmp = df[cols_encontradas].copy()
        tmp.columns = labels

        # Formatear columnas de dinero
        for label in labels:
            if label in ("VALOR", "MONTO"):
                tmp[label] = tmp[label].apply(lambda x: _fmt_dinero(x))

        filas = tmp.values.tolist()

        # Dividir en bloques según espacio disponible
        FILA_ALTO_APROX = 16
        while filas:
            espacio = self._espacio_disponible
            max_filas = max(1, int((espacio - 30) / FILA_ALTO_APROX))
            bloque = filas[:max_filas]
            filas = filas[max_filas:]

            data = [labels] + bloque
            tabla = Table(data, colWidths=anchos)

            estilos = [
                ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 8),
                ('BACKGROUND', (0, 0), (-1, 0), COLOR_HEADER),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTSIZE', (0, 1), (-1, -1), 7.5),
                ('TEXTCOLOR', (0, 1), (-1, -1), COLOR_TEXTO),
                ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CCCCCC')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ]

            # Filas alternas
            for i in range(1, len(data)):
                if i % 2 == 0:
                    estilos.append(('BACKGROUND', (0, i), (-1, i), COLOR_FILA_ALT))

            # Alineaciones
            for j, align in enumerate(aligns):
                if align == 'RIGHT':
                    estilos.append(('ALIGN', (j, 0), (j, -1), 'RIGHT'))
                elif align == 'CENTER':
                    estilos.append(('ALIGN', (j, 0), (j, -1), 'CENTER'))

            tabla.setStyle(TableStyle(estilos))
            w, h = tabla.wrapOn(self.c, self.ancho, self.alto)
            tabla.drawOn(self.c, self.MARGEN + 2, self.y - h)
            self.y -= (h + 8)

            if filas:
                self._nueva_pagina()

        self.y -= 10

    def clientes_sin_pago(self, lista):
        """Dibuja la lista de clientes sin pago."""
        self.seccion("CLIENTES SIN PAGO")

        if not lista:
            self.c.setFillColor(COLOR_GRIS)
            self.c.setFont('Helvetica-Oblique', 9)
            self.c.drawString(self.MARGEN + 10, self.y, "Todos los clientes al día ✓")
            self.y -= 20
            return

        # Convertir lista a tabla
        data = [["#", "CÓDIGO", "CLIENTE", "DÍAS DE ATRASO"]]
        for idx, item in enumerate(lista, 1):
            if isinstance(item, dict):
                cod = item['cod']
                nombre = item['nombre']
                dias = item['dias_atraso']
            else:
                partes = item.split()
                cod = partes[0] if len(partes) > 0 else ""
                nombre = partes[1] if len(partes) > 1 else ""
                try:
                    dias = int(partes[-4])
                except:
                    dias = 0
            data.append([str(idx), cod, nombre, str(dias)])

        ancho_util = self.ancho - 2 * self.MARGEN
        col_widths = [25, 55, ancho_util * 0.60, ancho_util * 0.25]

        FILA_ALTO = 15
        while len(data) > 1:
            espacio = self._espacio_disponible
            max_filas = max(1, int((espacio - 30) / FILA_ALTO))
            bloque = data[1:max_filas + 1]
            data = [data[0]] + data[max_filas + 1:]

            tabla_data = [data[0]] + bloque if len(data) > 0 else [["#", "CÓDIGO", "CLIENTE", "DÍAS DE ATRASO"]] + bloque
            tabla = Table(tabla_data, colWidths=col_widths)

            estilos = [
                ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 8),
                ('BACKGROUND', (0, 0), (-1, 0), COLOR_HEADER),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTSIZE', (0, 1), (-1, -1), 7),
                ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CCCCCC')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
                ('ALIGN', (0, 0), (0, -1), 'CENTER'),
                ('ALIGN', (1, 0), (1, -1), 'CENTER'),
                ('ALIGN', (3, 0), (3, -1), 'CENTER'),
            ]

            COLOR_VERDE_CLARO = colors.HexColor('#D4EDDA')
            COLOR_VERDE_TEXTO = colors.HexColor('#155724')
            COLOR_AMARILLO_CLARO = colors.HexColor('#FFF3CD')
            COLOR_AMARILLO_TEXTO = colors.HexColor('#856404')
            COLOR_ROJO_CLARO = colors.HexColor('#F8D7DA')
            COLOR_ROJO_TEXTO = colors.HexColor('#721C24')

            for i in range(1, len(tabla_data)):
                try:
                    dias = int(tabla_data[i][3])
                    if dias <= 2:
                        bg_color = COLOR_VERDE_CLARO
                        txt_color = COLOR_VERDE_TEXTO
                    elif 3 <= dias <= 5:
                        bg_color = COLOR_AMARILLO_CLARO
                        txt_color = COLOR_AMARILLO_TEXTO
                    else:
                        bg_color = COLOR_ROJO_CLARO
                        txt_color = COLOR_ROJO_TEXTO
                except:
                    bg_color = COLOR_FILA_ALT if i % 2 == 0 else colors.white
                    txt_color = COLOR_TEXTO
                    
                estilos.append(('BACKGROUND', (0, i), (-1, i), bg_color))
                estilos.append(('TEXTCOLOR', (0, i), (-1, i), txt_color))

            tabla.setStyle(TableStyle(estilos))
            w, h = tabla.wrapOn(self.c, self.ancho, self.alto)
            tabla.drawOn(self.c, self.MARGEN + 2, self.y - h)
            self.y -= (h + 8)

            if len(data) > 1:
                self._nueva_pagina()

    def guardar(self):
        """Guarda el PDF."""
        self._dibujar_footer()
        self.c.save()


# ============================================================
# FUNCIÓN PRINCIPAL
# ============================================================

def main():
    # Fecha del día anterior
    fecha_base = datetime.now() - timedelta(days=1)
    fecha_informe = fecha_base.strftime("%d/%m/%Y")
    fecha_archivo = fecha_base.strftime("%d-%m-%Y")

    print("=" * 60)
    print(f"  GENERADOR DE REPORTES POR RUTA")
    print(f"  Fecha del informe: {fecha_informe} (ayer)")
    print(f"  Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    print("=" * 60)

    # Validar carpeta de entradas
    if not os.path.isdir(CARPETA_ENTRADAS):
        print(f"\n❌ No se encontró la carpeta de entradas: {CARPETA_ENTRADAS}")
        sys.exit(1)

    archivos_entrada = os.listdir(CARPETA_ENTRADAS)
    if not archivos_entrada:
        print(f"\n❌ La carpeta de entradas está vacía: {CARPETA_ENTRADAS}")
        sys.exit(1)

    print(f"\n📁 Archivos en entradas/: {len(archivos_entrada)}")
    for f in sorted(archivos_entrada):
        print(f"   • {f}")

    # Crear carpeta de salida
    os.makedirs(CARPETA_SALIDAS, exist_ok=True)

    # ── 1. CUADRE ──
    print(f"\n{'─' * 40}")
    print("📊 Leyendo CUADRE...")
    archivos_cuadre = encontrar_archivos("ReporteExcelCuadreVariasRutas")
    datos_cuadre = {}
    for arch in archivos_cuadre:
        print(f"   ← {os.path.basename(arch)}")
        datos_cuadre.update(leer_cuadre_por_rutas(arch))
    print(f"   ✓ {len(datos_cuadre)} rutas encontradas")

    # ── 2. GASTOS ──
    print(f"\n{'─' * 40}")
    print("💰 Leyendo GASTOS...")
    archivos_gastos = encontrar_archivos("ReporteExcelGastos")
    lista_gastos = []
    for arch in archivos_gastos:
        print(f"   ← {os.path.basename(arch)}")
        df = leer_excel_tabla(arch)
        if df is not None:
            lista_gastos.append(df)
            print(f"     {df.shape[0]} filas")
    df_gastos = pd.concat(lista_gastos, ignore_index=True) if lista_gastos else pd.DataFrame()

    # ── 3. CRÉDITOS ──
    print(f"\n{'─' * 40}")
    print("📋 Leyendo CRÉDITOS...")
    archivos_creditos = encontrar_archivos("ReporteExcelCreditos")
    lista_creditos = []
    for arch in archivos_creditos:
        print(f"   ← {os.path.basename(arch)}")
        df = leer_excel_tabla(arch)
        if df is not None:
            lista_creditos.append(df)
            print(f"     {df.shape[0]} filas")
    df_creditos = pd.concat(lista_creditos, ignore_index=True) if lista_creditos else pd.DataFrame()

    # ── 4. CLIENTES SIN PAGO ──
    print(f"\n{'─' * 40}")
    print("🔴 Leyendo CLIENTES SIN PAGO...")
    archivos_pdf = encontrar_archivos("ReportePdfClientesSinPagos", extension=".pdf")
    clientes_combinados = {}
    for arch in archivos_pdf:
        print(f"   ← {os.path.basename(arch)}")
        datos = extraer_clientes_pdf(arch)
        for ruta, cls in datos.items():
            clientes_combinados.setdefault(ruta, []).extend(cls)
    total_clientes = sum(len(v) for v in clientes_combinados.values())
    print(f"   ✓ {len(clientes_combinados)} rutas, {total_clientes} clientes")

    # ── 5. GENERAR PDFs ──
    print(f"\n{'=' * 60}")
    print("📄 GENERANDO REPORTES PDF...")
    print(f"{'=' * 60}")

    generados = 0
    errores = 0

    rutas_encontradas = set(MIS_RUTAS)
    rutas_encontradas.update(datos_cuadre.keys())
    if not df_gastos.empty:
        col = _find_col(df_gastos, ["RUTA"])
        if col:
            rutas_encontradas.update(df_gastos[col].dropna().astype(str).str.strip().unique())
    if not df_creditos.empty:
        col = _find_col(df_creditos, ["RUTA"])
        if col:
            rutas_encontradas.update(df_creditos[col].dropna().astype(str).str.strip().unique())
    rutas_encontradas.update(clientes_combinados.keys())

    todas_las_rutas = sorted([r for r in rutas_encontradas if r])

    for ruta in todas_las_rutas:
        try:
            cuadre = datos_cuadre.get(ruta, {})
            gastos_sub = filtrar_por_ruta(df_gastos, ruta)
            creditos_sub = filtrar_por_ruta(df_creditos, ruta)
            clientes = clientes_combinados.get(ruta, [])

            # Ignorar rutas que no tienen ningún dato si no están en MIS_RUTAS
            if not cuadre and gastos_sub.empty and creditos_sub.empty and not clientes and ruta not in MIS_RUTAS:
                continue

            nombre_limpio = ruta.replace(".", "").replace(" ", "_").replace("/", "_")
            nombre_pdf = f"Reporte_{nombre_limpio}_{fecha_archivo}.pdf"
            salida = os.path.join(CARPETA_SALIDAS, nombre_pdf)

            # Crear PDF
            pdf = ReportePDF(salida, ruta, fecha_informe)

            pdf.cuadre(cuadre)

            pdf.tabla("GASTOS Y MOVIMIENTOS", gastos_sub, [
                (["OBSERVACION", "OBSERVACIÓN", "CONCEPTO"], "DESCRIPCIÓN", 0.65, 'LEFT'),
                (["VALOR", "CAUSANTES"], "VALOR", 0.25, 'RIGHT'),
            ])

            pdf.tabla("CRÉDITOS", creditos_sub, [
                (["NOMBRE", "CLIENTE"], "CLIENTE", 0.42, 'LEFT'),
                (["TELEFONO", "TELÉFONO", "CELULAR"], "TELÉFONO", 0.28, 'LEFT'),
                (["VALOR CRÉDITO", "VALOR CREDITO", "VALOR"], "VALOR", 0.20, 'RIGHT'),
            ])

            pdf.clientes_sin_pago(clientes)
            pdf.guardar()

            generados += 1
            marca_ok = "✓" if cuadre else "○"
            print(f"  {marca_ok} {ruta:35s} → {nombre_pdf}")
            print(f"     Cuadre: {'Sí' if cuadre else 'No'} | "
                  f"Gastos: {len(gastos_sub)} | "
                  f"Créditos: {len(creditos_sub)} | "
                  f"Sin pago: {len(clientes)}")

        except Exception as e:
            errores += 1
            print(f"  ❌ {ruta}: {e}")

    # ── RESUMEN FINAL ──
    print(f"\n{'=' * 60}")
    print(f"  ✅ COMPLETADO")
    print(f"  PDFs generados: {generados}/{len(todas_las_rutas)}")
    if errores:
        print(f"  ⚠  Errores: {errores}")
    print(f"  📁 Carpeta: {CARPETA_SALIDAS}/")
    print(f"{'=' * 60}")


if __name__ == "__main__":
    main()
